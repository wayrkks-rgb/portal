"""통합기 할당률과 메모리 표기를 확인한다.

사용률과 할당률은 다르다. 사용률은 '지금 실제로 쓰는 양' 이고 할당률은 'VM 에게
나눠준 양' 이다. 메모리 1TB 짜리 통합기에 4GB 씩 10 대를 만들었다면 실제 사용률이
5% 라도 할당률은 40/1024 이다. VM 을 더 만들 수 있는지는 할당률로만 알 수 있으므로
두 값을 함께 낸다.

메모리 표기는 두 가지를 지킨다.

1. ESXi 는 하이퍼바이저가 쓰는 몫을 뺀 값을 알려준다. 1TB 장비가 1023.66GB 로
   나오므로 장표에 적을 때는 반올림해 1024 로 적는다.
2. 엑셀 표시형식 ``0.##`` 은 16 을 "16." 으로 그린다. 소수 자리가 비어도 점은
   찍기 때문이다. 장표에 점이 남으면 안 된다.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from asset_sync.config import AppConfig
from asset_sync.db.sqlite_manager import SQLiteManager
from asset_sync.repositories import AssetRepository
from asset_sync.services.resource_usage_service import VMResourceUsageExportService

STAT_DATE = "2026-09-01"
GB = 1024


def _service(tmp_path: Path):
    cfg = AppConfig(
        root_dir=tmp_path,
        sqlite_path=Path("data/alloc.db"),
        rvtools={"resource_usage": {"enabled": False}, "vcenters": []},
    )
    manager = SQLiteManager(cfg.database_path)
    manager.initialize()
    return cfg, manager


def _seed(conn, hosts: list[dict], vms: list[dict]) -> None:
    """통합기(ESXi) 용량과 그 위 VM 할당량을 넣는다."""
    conn.execute(
        "INSERT INTO resource_usage_run(period_start, period_end, started_at, status)"
        " VALUES(?, ?, ?, 'SUCCESS')",
        (STAT_DATE, STAT_DATE, STAT_DATE),
    )
    run_id = 1
    for host in hosts:
        conn.execute(
            "INSERT INTO host_resource_usage_daily(run_id, stat_date, vcenter_id, service_name,"
            " cluster_name, esxi_host, vm_count, allocated_cpu_cores, allocated_memory_mb,"
            " cpu_max_pct, cpu_avg_pct, mem_max_pct, mem_avg_pct, sample_count, collection_status,"
            " raw_json, created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'SUCCESS', '{}', ?)",
            (run_id, STAT_DATE, host["vcenter_id"], host["service_name"], host["cluster_name"],
             host["esxi_host"], host["vm_count"], host["cpu"], host["memory_mb"],
             80.0, 40.0, 70.0, 35.0, 12, STAT_DATE),
        )
    for vm in vms:
        conn.execute(
            "INSERT INTO vm_resource_usage_daily(run_id, stat_date, vcenter_snapshot_id, asset_key,"
            " vcenter_id, service_name, cluster_name, esxi_host, vm_uuid, vm_name, power_state,"
            " allocated_cpu_cores, allocated_memory_mb, cpu_max_pct, cpu_avg_pct, mem_max_pct,"
            " mem_avg_pct, sample_count, inventory_status, collection_status, raw_json, created_at)"
            "VALUES(?,?,NULL,?,?,?,?,?,?,?, 'poweredOn', ?,?, 30.0, 10.0, 60.0, 30.0, 12, ?, 'SUCCESS', '{}', ?)",
            (run_id, STAT_DATE, vm["vm_name"], vm["vcenter_id"], vm["service_name"], vm["cluster_name"],
             vm["esxi_host"], vm["vm_name"], vm["vm_name"], vm["cpu"], vm["memory_mb"],
             vm.get("inventory_status", "CURRENT"), STAT_DATE),
        )
    conn.commit()


def test_a_terabyte_host_is_written_as_1024_not_1023_66(tmp_path: Path) -> None:
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 1, "cpu": 64, "memory_mb": 1048234,
        }], [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_name": "vm-1", "cpu": 4, "memory_mb": 4 * GB,
        }])
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    assert result["hosts"][0]["allocated_memory_gb"] == 1024
    # 512MB 짜리는 0 이 되면 안 된다.
    assert VMResourceUsageExportService._mb_to_gb(512) == 0.5


def test_the_allocation_rate_divides_what_vms_got_by_what_the_host_has(tmp_path: Path) -> None:
    """양식 그대로: 1TB 통합기에 4GB 씩 10 대면 MEM 할당률은 40/1024 이다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 10, "cpu": 64, "memory_mb": 1024 * GB,
        }], [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_name": f"vm-{index}", "cpu": 2, "memory_mb": 4 * GB,
        } for index in range(10)])
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    host = result["hosts"][0]
    assert host["assigned_memory_gb"] == 40
    assert host["mem_alloc_pct"] == round(40 / 1024 * 100, 2)
    assert host["assigned_cpu_cores"] == 20
    assert host["cpu_alloc_pct"] == round(20 / 64 * 100, 2)
    # 사용률은 그대로 남는다. 둘 다 있어야 판단할 수 있다.
    assert host["mem_avg_pct"] == 35.0


def test_a_deleted_vm_no_longer_counts_toward_allocation(tmp_path: Path) -> None:
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 1, "cpu": 10, "memory_mb": 100 * GB,
        }], [
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-01", "vm_name": "vm-live", "cpu": 2, "memory_mb": 8 * GB},
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-01", "vm_name": "vm-gone", "cpu": 8, "memory_mb": 64 * GB,
             "inventory_status": "NOT_IN_CURRENT_INVENTORY"},
        ])
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    assert result["hosts"][0]["assigned_cpu_cores"] == 2
    assert result["hosts"][0]["assigned_memory_gb"] == 8


def test_a_cluster_row_sums_the_esxi_hosts_under_it(tmp_path: Path) -> None:
    """통합기 한 대는 ESXi 여러 대의 묶음이다. 남은 여유는 묶음 단위로 봐야 한다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-01", "vm_count": 1, "cpu": 32, "memory_mb": 512 * GB},
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-02", "vm_count": 1, "cpu": 32, "memory_mb": 512 * GB},
        ], [
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-01", "vm_name": "vm-1", "cpu": 8, "memory_mb": 64 * GB},
            {"vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
             "esxi_host": "esxi-02", "vm_name": "vm-2", "cpu": 8, "memory_mb": 64 * GB},
        ])
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    assert len(result["clusters"]) == 1
    cluster = result["clusters"][0]
    assert cluster["host_count"] == 2
    assert cluster["allocated_memory_gb"] == 1024
    assert cluster["assigned_memory_gb"] == 128
    assert cluster["cpu_alloc_pct"] == 25.0
    assert result["summary"]["cluster_count"] == 1


def test_a_host_with_no_capacity_reported_has_no_allocation_rate(tmp_path: Path) -> None:
    """용량을 모르면 0% 가 아니라 '모름' 이다. 0% 로 적으면 여유가 있다고 읽힌다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 1, "cpu": None, "memory_mb": None,
        }], [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_name": "vm-1", "cpu": 4, "memory_mb": 8 * GB,
        }])
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    assert result["hosts"][0]["cpu_alloc_pct"] is None
    assert result["hosts"][0]["mem_alloc_pct"] is None


def test_the_excel_never_leaves_a_trailing_dot_behind_a_whole_number(tmp_path: Path) -> None:
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        _seed(conn, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 1, "cpu": 64, "memory_mb": 1048234,
        }], [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_name": "vm-1", "cpu": 4, "memory_mb": 16 * GB,
        }])
        target = VMResourceUsageExportService(cfg, AssetRepository(conn)).export_xlsx(
            STAT_DATE, STAT_DATE, tmp_path / "export"
        )

    workbook = load_workbook(target)
    sheet = workbook["VMsResource"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("실제 Memory GB") + 1
    cell = sheet.cell(row=2, column=column)
    assert cell.value == 16
    assert "0.##" not in str(cell.number_format), "0.## 은 16 을 '16.' 으로 그린다"
    workbook.close()
