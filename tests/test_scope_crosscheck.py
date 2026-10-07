"""VM 제외와 서버현황이 어긋난 서버를 찾아내야 한다.

VM 을 전원 꺼짐 등으로 자원 집계에서 뺐는데 같은 서버가 ITSM 서버현황에는 자산
으로 남아 있다면 둘 중 하나가 틀렸다. 어느 쪽이 맞는지는 사람이 알지만, **어긋난
것을 찾아 이유와 함께 보여주는 것**은 기계가 해야 한다.

반대 방향도 같다. ITSM 은 미사용인데 VM 은 켜져 돌고 있으면 실물은 살아 있고
장부만 죽은 것이다.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asset_sync.config import AppConfig
from asset_sync.db.manager import create_manager
from asset_sync.repositories import AssetRepository
from asset_sync.services import asset_matching
from asset_sync.services.scope_crosscheck_service import (
    AGREED,
    AMBIGUOUS,
    BOTH_EXCLUDED,
    ITSM_EXCLUDED,
    NO_VM,
    POWERED_OFF,
    VM_EXCLUDED,
    VM_ONLY,
    ScopeCrossCheckService,
)

ROOT = Path(__file__).resolve().parents[1]


def asset(cm_id: str, *, hostname: str, ip: str, status: str = "CMSTA010",
          physical: bool = False) -> dict:
    return {
        "cm_id": cm_id,
        "normalized_hostname": hostname,
        "primary_ip": ip,
        "ip_addresses": [ip],
        "os_family": "Linux Redhat", "os_version": "8.6", "status_code": status,
        "server_category_code": "CMSVRCATCD010" if physical else "CMSVRCATCD020",
        "raw": {
            "CM_ID": cm_id, "CM_NAME": f"업무-{cm_id}", "CM_HOSTNAME": hostname,
            "CM_IP": ip, "CM_OS": "CMCIOSCD010", "CM_OS_VERSION": "8.6",
            "CM_EOL_DT": "2030-12-31", "CM_PLACE": "CMPLACE010",
        },
    }


def vm(name: str, *, hostname: str, ip: str, power: str = "poweredOn",
       template: bool = False) -> dict:
    return {
        "asset_key": name, "vm_uuid": f"uuid-{name}", "smbios_uuid": f"sm-{name}",
        "vm_id": f"mo-{name}", "vcenter": "VC1", "vm_name": name,
        "dns_name": hostname, "normalized_hostname": hostname,
        "ip_addresses": [ip], "primary_ip": ip, "cpus": 4, "memory_mb": 8192,
        "os_family": "Linux Redhat", "os_version": "8.6", "power_state": power,
        "datacenter": "DC1", "cluster": "CL1", "cluster_name": "CL1",
        "esxi_host": "esxi-01", "template_flag": template, "srm_placeholder": False,
        "raw": {}, "record_hash": "h",
    }


@pytest.fixture()
def portal(tmp_path: Path):
    config = AppConfig(root_dir=tmp_path, sqlite_path=Path("data/cross.db"))
    manager = create_manager(config)
    manager.initialize()
    return config, manager


def seed(manager, assets: list[dict], vms: list[dict]) -> None:
    now = datetime.now()
    with manager.connect() as conn:
        repo = AssetRepository(conn)
        run = repo.start_collection_run("ITSM", now.isoformat())
        repo.finish_collection_run(run, "SUCCESS", len(assets), now.isoformat())
        snapshot = repo.create_snapshot(
            "ITSM", now.date().isoformat(), now.isoformat(), run, "SUCCESS", len(assets), "h"
        )
        conn.executemany(
            "INSERT INTO itsm_asset_snapshot(snapshot_id,cm_id,normalized_hostname,primary_ip,"
            "ip_json,cpu_cores,memory_mb,os_family,os_version,status_code,server_category_code,"
            "environment_code,eos_value,record_hash,raw_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(snapshot, r["cm_id"], r["normalized_hostname"], r["primary_ip"],
              json.dumps(r["ip_addresses"]), 4, 8192, r["os_family"], r["os_version"],
              r["status_code"], r["server_category_code"], "CMOWNCATCD0010",
              r["raw"]["CM_EOL_DT"], "h", json.dumps(r["raw"], ensure_ascii=False))
             for r in assets],
        )
        rv_run = repo.start_collection_run("RVTOOLS", now.isoformat())
        repo.finish_collection_run(rv_run, "SUCCESS", len(vms), now.isoformat(), ["VC1"], [])
        rv_snapshot = repo.create_snapshot(
            "RVTOOLS", now.date().isoformat(), now.isoformat(), rv_run, "SUCCESS", len(vms), "h"
        )
        repo.insert_rv_records(rv_snapshot, vms)
        conn.commit()


def exclude_vm(manager, asset_key: str, reason: str = "전원 꺼짐") -> None:
    from asset_sync.services import asset_scope

    with manager.connect() as conn:
        asset_scope.save_rules(AssetRepository(conn), "RVTOOLS", [{"asset_key": asset_key}],
                               mode="EXCLUDE", reason=reason)
        conn.commit()


def check(config, manager) -> dict:
    with manager.connect() as conn:
        return ScopeCrossCheckService(config, AssetRepository(conn)).check()


def verdicts(result: dict) -> dict[str, str]:
    return {str(item.get("cm_id") or item.get("vm_name")): item["verdict"]
            for item in result["items"]}


# ── 핵심: 어긋난 것을 찾는다 ───────────────────────────────────────────


def test_a_vm_excluded_while_itsm_still_counts_it_is_flagged(portal) -> None:
    """이게 물어보신 바로 그 경우다. VM 은 뺐는데 자산으로 세고 있다."""
    config, manager = portal
    seed(manager,
         [asset("CM001", hostname="srv-a", ip="10.0.0.1"),
          asset("CM002", hostname="srv-b", ip="10.0.0.2")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1"),
          vm("vm-b", hostname="srv-b", ip="10.0.0.2")])
    exclude_vm(manager, "vm-a", "전원 꺼진 상태로 6개월")

    result = check(config, manager)
    assert verdicts(result)["CM001"] == VM_EXCLUDED
    assert verdicts(result)["CM002"] == AGREED

    row = next(item for item in result["items"] if item["cm_id"] == "CM001")
    assert row["review"] is True
    assert "VM 은 뺐는데" in row["verdict_label"]
    # 어느 쪽을 고쳐야 하는지 알려줘야 한다.
    assert "ITSM 상태를" in row["action"] and "되돌리" in row["action"]
    # 양쪽 값이 다 있어야 그 자리에서 판단할 수 있다.
    assert row["hostname"] == "srv-a" and row["primary_ip"] == "10.0.0.1"
    assert row["vm_name"] == "vm-a"
    assert "제외" in row["why"] and "자산" in row["why"]
    assert result["summary"]["review"] == 1


def test_a_powered_off_vm_counted_as_an_asset_is_flagged(portal) -> None:
    """수동으로 빼지 않았어도 전원이 꺼져 있으면 확인 대상이다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1", power="poweredOff")])

    result = check(config, manager)
    row = result["items"][0]
    assert row["verdict"] == POWERED_OFF
    assert row["review"] is True
    assert row["power_state"] == "poweredOff"
    assert row["powered_off"] is True


def test_a_template_vm_counted_as_an_asset_is_flagged(portal) -> None:
    """템플릿도 자원 집계에서 빠진다. 자산으로 세고 있으면 어긋난 것이다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1", template=True)])
    row = check(config, manager)["items"][0]
    assert row["verdict"] == VM_EXCLUDED
    assert row["vm_reason"] == "TEMPLATE"
    assert row["vm_reason_label"], "사유를 사람이 읽을 수 있게 적어야 한다"


def test_the_other_direction_is_flagged_too(portal) -> None:
    """ITSM 은 미사용인데 VM 은 켜져 돌고 있다. 장부만 죽은 것이다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1", status="CMSTA020")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1")])

    row = check(config, manager)["items"][0]
    assert row["verdict"] == ITSM_EXCLUDED
    assert row["review"] is True
    assert row["asset_included"] is False
    assert "실물이 돌고 있습니다" in row["action"]


def test_both_excluded_is_not_a_problem(portal) -> None:
    """정리를 끝낸 것은 확인 대상이 아니다. 목록이 쓸모없어진다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1", status="CMSTA060")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1", power="poweredOff")])
    exclude_vm(manager, "vm-a")

    row = check(config, manager)["items"][0]
    assert row["verdict"] == BOTH_EXCLUDED
    assert row["review"] is False


def test_a_logical_asset_without_any_vm_is_flagged(portal) -> None:
    """논리서버인데 VM 이 없다. 사라졌거나 짝을 못 찾은 것이다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1")], [])
    row = check(config, manager)["items"][0]
    assert row["verdict"] == NO_VM
    assert row["review"] is True
    assert "VM 없음" in row["why"]


def test_a_physical_asset_without_a_vm_is_normal(portal) -> None:
    """물리서버는 VM 이 없는 것이 당연하다. 올리면 목록이 쓸모없어진다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1", physical=True)], [])
    row = check(config, manager)["items"][0]
    assert row["verdict"] == AGREED
    assert row["review"] is False


def test_a_vm_without_any_asset_is_listed_but_not_review(portal) -> None:
    config, manager = portal
    seed(manager, [], [vm("vm-x", hostname="srv-x", ip="10.0.0.9")])
    row = check(config, manager)["items"][0]
    assert row["verdict"] == VM_ONLY
    assert row["vm_name"] == "vm-x"


def test_system_vms_are_not_listed_as_orphans(portal) -> None:
    """vCLS 까지 'ITSM 에 없는 VM' 으로 올리면 목록을 못 쓴다."""
    config, manager = portal
    seed(manager, [], [vm("vCLS-1234", hostname="", ip="10.0.0.8"),
                       vm("vm-x", hostname="srv-x", ip="10.0.0.9")])
    names = {item["vm_name"] for item in check(config, manager)["items"]}
    assert names == {"vm-x"}


def test_two_candidates_are_not_paired_blindly(portal) -> None:
    """틀린 짝을 지으면 양쪽 대수가 다 틀린다. 짝으로 보지 않는다."""
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1")],
         [vm("vm-a1", hostname="srv-a", ip="10.0.0.1"),
          vm("vm-a2", hostname="srv-a", ip="10.0.0.1")])
    row = next(item for item in check(config, manager)["items"] if item["cm_id"] == "CM001")
    assert row["verdict"] == AMBIGUOUS
    assert len(row["candidates"]) == 2
    assert "정합성 화면" in row["action"]


def test_the_counts_add_up(portal) -> None:
    config, manager = portal
    seed(manager,
         [asset("CM001", hostname="srv-a", ip="10.0.0.1"),
          asset("CM002", hostname="srv-b", ip="10.0.0.2"),
          asset("CM003", hostname="srv-c", ip="10.0.0.3", status="CMSTA020")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1"),
          vm("vm-b", hostname="srv-b", ip="10.0.0.2", power="poweredOff"),
          vm("vm-c", hostname="srv-c", ip="10.0.0.3")])
    exclude_vm(manager, "vm-a")

    result = check(config, manager)
    assert result["summary"]["checked"] == 3
    assert result["summary"]["review"] == 3
    counts = result["counts"]
    assert counts[VM_EXCLUDED]["count"] == 1
    assert counts[POWERED_OFF]["count"] == 1
    assert counts[ITSM_EXCLUDED]["count"] == 1
    # 합이 맞아야 한다. 어딘가 빠지면 '확인 필요' 를 놓친다.
    assert sum(item["count"] for item in counts.values()) == len(result["items"])


def test_without_both_snapshots_it_says_so(portal) -> None:
    config, manager = portal
    result = check(config, manager)
    assert result["status"] == "NO_SNAPSHOT"
    assert "모두 있어야" in result["message"]


# ── 짝짓기 규칙은 한 곳에만 둔다 ───────────────────────────────────────


def test_the_matcher_prefers_the_remembered_pair() -> None:
    records = {"vm-a": {"vm_uuid": "u1", "normalized_hostname": "other", "ip_addresses": []}}
    index = asset_matching.build_index(records)
    match = asset_matching.find(
        "CM001", {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]},
        index, identity_map={"CM001": "u1"},
    )
    assert match.key == "vm-a" and match.method == "IDENTITY_MAP" and match.score == 100


def test_the_matcher_ranks_hostname_and_ip_above_either_alone() -> None:
    records = {
        "both": {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]},
        "host": {"normalized_hostname": "srv-a", "ip_addresses": ["10.9.9.9"]},
    }
    index = asset_matching.build_index(records)
    match = asset_matching.find(
        "CM001", {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]}, index
    )
    assert match.key == "both" and match.method == "IP_HOSTNAME"


def test_the_matcher_reports_ambiguity_instead_of_guessing() -> None:
    records = {
        "a": {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]},
        "b": {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]},
    }
    index = asset_matching.build_index(records)
    match = asset_matching.find(
        "CM001", {"normalized_hostname": "srv-a", "ip_addresses": ["10.0.0.1"]}, index
    )
    assert match.key is None and match.ambiguous is True
    assert match.candidates == ("a", "b")


def test_reconciliation_uses_the_same_matcher() -> None:
    """두 화면이 다른 규칙으로 짝을 지으면 어느 쪽을 믿어야 할지 알 수 없다."""
    source = (ROOT / "asset_sync" / "services" / "reconciliation_service.py").read_text(
        encoding="utf-8")
    assert "asset_matching.build_index" in source
    assert "asset_matching.find" in source
    # 예전 인라인 색인이 되살아나지 않게 못 박는다.
    assert "host_index" not in source and "ip_index" not in source


# ── 화면에서 쓸 수 있어야 한다 ─────────────────────────────────────────


@pytest.fixture()
def web(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    (tmp_path / "config").mkdir()
    (tmp_path / "config" / "app_config.yaml").write_text(
        "itsm:\n  collection_mode: DEMO\n", encoding="utf-8")
    monkeypatch.setenv("ASSET_APP_ROOT", str(tmp_path))
    monkeypatch.setenv("FLASK_SECRET_KEY", "test-secret")
    monkeypatch.chdir(tmp_path)
    from application import create_app
    from application.db import reset_database_manager

    reset_database_manager()
    yield create_app()
    reset_database_manager()


def test_the_screen_gets_only_what_needs_review_by_default(web) -> None:
    from asset_sync.config import load_config
    from asset_sync.db.manager import create_manager as make

    config = load_config()
    manager = make(config)
    seed(manager,
         [asset("CM001", hostname="srv-a", ip="10.0.0.1"),
          asset("CM002", hostname="srv-b", ip="10.0.0.2")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1", power="poweredOff"),
          vm("vm-b", hostname="srv-b", ip="10.0.0.2")])

    client = web.test_client()
    with client.session_transaction() as session:
        session["user"] = {"id": 1, "username": "admin", "role": "admin", "name": "admin"}

    payload = client.get("/api/asset-sync/scope-crosscheck").get_json()
    assert [item["cm_id"] for item in payload["items"]] == ["CM001"]
    assert payload["summary"]["review"] == 1

    everything = client.get("/api/asset-sync/scope-crosscheck?review=all").get_json()
    assert len(everything["items"]) == 2

    response = client.get("/api/asset-sync/scope-crosscheck/export")
    assert response.status_code == 200
    assert response.data[:2] == b"PK"


def test_the_export_puts_the_verdict_and_the_action_first(portal, tmp_path: Path) -> None:
    config, manager = portal
    seed(manager, [asset("CM001", hostname="srv-a", ip="10.0.0.1")],
         [vm("vm-a", hostname="srv-a", ip="10.0.0.1", power="poweredOff")])
    result = check(config, manager)
    path = tmp_path / "cross.xlsx"
    ScopeCrossCheckService.write_xlsx(result, result["items"], path)

    sheet = load_workbook(path).active
    head = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "판정")
    labels = [sheet.cell(head, column).value for column in range(1, 6)]
    assert labels[:3] == ["판정", "해야 할 일", "왜"]
    assert "자산번호" in labels and "호스트명" in labels
    assert sheet.cell(head + 1, 1).value == "자산인데 VM 전원이 꺼져 있음"
    assert "확인 필요" in str(sheet["A2"].value)


def test_the_card_is_wired_on_the_screen() -> None:
    """양쪽 id 가 맞아야 그려진다. 한쪽만 고치면 조용히 빈 화면이 된다."""
    page = (ROOT / "templates" / "pages" / "monthly_check.html").read_text(encoding="utf-8")
    script = (ROOT / "templates" / "partials" / "js" / "monthly_check.html").read_text(
        encoding="utf-8")
    for element_id in ("crosscheck-table", "crosscheck-counts", "crosscheck-alert",
                       "crosscheck-basis", "crosscheck-search", "crosscheck-all"):
        assert f'id="{element_id}"' in page, f"{element_id} 가 화면에 없다"
        assert f"'{element_id}'" in script, f"{element_id} 를 스크립트가 안 찾는다"
    assert "/api/asset-sync/scope-crosscheck" in script
    assert "loadCrossCheck" in script
    # 그 자리에서 자산을 뺄 수 있어야 한다.
    assert "excludeOneAsset" in script.split("function renderCrossCheckAgain", 1)[1]
