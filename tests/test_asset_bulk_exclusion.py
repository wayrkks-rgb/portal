"""엑셀로 수백 건을 한꺼번에 제외하는 길.

화면에서 300 건을 하나씩 체크할 수는 없다. 목록을 엑셀로 받아 [처리] 칸에 적고
그대로 다시 올리면 한 번에 적용한다. 받은 파일을 그대로 되올릴 수 있어야 하므로
**내보낸 파일과 읽는 쪽이 맞물리는지**를 여기서 지킨다.

조용히 넘기지 않는 것도 지킨다. 300 줄 중 몇 줄이 오타라서 빠졌다는 것을 모르면
숫자가 틀린 채로 보고가 나간다.
"""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asset_sync.config import AppConfig
from asset_sync.db.manager import create_manager
from asset_sync.repositories import AssetRepository
from asset_sync.services.asset_scope import (
    ACTION_CHOICES,
    AssetScopeError,
    apply_bulk,
    list_rules,
    read_bulk_sheet,
)
from asset_sync.services.monthly_export_service import MonthlyCheckExportService
from asset_sync.services.server_status_service import ServerStatusService

ASSETS = 300


def raw(cm_id: str) -> dict:
    return {
        "CM_ID": cm_id, "CM_NAME": f"업무-{cm_id}", "CM_HOSTNAME": f"host-{cm_id}",
        "CM_IP": "10.0.0.5", "CM_OS": "CMCIOSCD010", "CM_OS_VERSION": "8.6",
        "CM_EOL_DT": "2030-12-31", "CM_PLACE": "CMPLACE010",
    }


@pytest.fixture()
def portal(tmp_path: Path):
    config = AppConfig(root_dir=tmp_path, sqlite_path=Path("data/bulk.db"))
    manager = create_manager(config)
    manager.initialize()
    now = datetime.now()
    rows = [f"CM{index:04d}" for index in range(ASSETS)]
    with manager.connect() as conn:
        repo = AssetRepository(conn)
        run = repo.start_collection_run("ITSM", now.isoformat())
        repo.finish_collection_run(run, "SUCCESS", len(rows), now.isoformat(), ["ALL"])
        snapshot_id = repo.create_snapshot(
            "ITSM", now.date().isoformat(), now.isoformat(), run, "SUCCESS", len(rows), "h")
        conn.executemany(
            "INSERT INTO itsm_asset_snapshot(snapshot_id,cm_id,normalized_hostname,primary_ip,"
            "ip_json,cpu_cores,memory_mb,os_family,os_version,status_code,server_category_code,"
            "environment_code,eos_value,record_hash,raw_json) VALUES(?,?,?,?,'[]',?,?,?,?,?,?,?,?,?,?)",
            [(snapshot_id, cm_id, f"host-{cm_id}", "10.0.0.5", 4, 8192, "Linux Redhat", "8.6",
              "CMSTA010", "CMSVRCATCD020", "CMOWNCATCD0010", "2030-12-31", "h",
              json.dumps(raw(cm_id), ensure_ascii=False)) for cm_id in rows],
        )
        conn.commit()
    return config, manager, snapshot_id


def export_list(config, manager, snapshot_id, target: Path) -> Path:
    with manager.connect() as conn:
        service = ServerStatusService(config, AssetRepository(conn), include_all=True)
        records = service.records(snapshot_id)
    workbook = MonthlyCheckExportService({"as_of": "2026-10-01"}, records).build_asset_list(
        "자산 목록", records, "조건: 전체")
    path = target / "자산목록.xlsx"
    target.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def key_column(sheet) -> int:
    """읽는 쪽이 쓰는 자산번호 열. 머리글 윗줄(한글 이름)이 먼저 걸린다."""
    labels = [cell.value for cell in sheet[4]]
    return labels.index("자산번호") + 1


def mark(path: Path, how: dict[str, str], *, reasons: dict[str, str] | None = None) -> Path:
    """사람이 엑셀에서 [처리] 칸을 채운 것과 같게 만든다."""
    workbook = load_workbook(path)
    sheet = workbook.active
    key_at = key_column(sheet)
    for line in range(6, sheet.max_row + 1):
        cm_id = sheet.cell(row=line, column=key_at).value
        if cm_id in how:
            sheet.cell(row=line, column=1, value=how[cm_id])
            if reasons and cm_id in reasons:
                sheet.cell(row=line, column=2, value=reasons[cm_id])
    workbook.save(path)
    workbook.close()
    return path


# ── 내보낸 파일이 되올릴 수 있는 모양인가 ─────────────────────────────────
def test_the_exported_sheet_has_a_column_to_fill_in(portal, tmp_path):
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    sheet = load_workbook(path).active

    labels = [cell.value for cell in sheet[4]]
    assert labels[0] == "처리"
    assert labels[1] == "제외 사유"
    # 쓰는 법이 파일 안에 적혀 있어야 설명을 따로 찾지 않는다.
    assert "처리" in str(sheet["A3"].value)
    assert "다시 올리면" in str(sheet["A3"].value)
    # 300 줄을 손으로 쓰면 오타가 난다. 고를 수 있어야 한다.
    validations = [v for v in sheet.data_validations.dataValidation]
    assert validations, "고르는 목록이 없습니다"
    assert all(word in validations[0].formula1 for word in ACTION_CHOICES)


def test_three_hundred_rows_go_in_at_once(portal, tmp_path):
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    wanted = {f"CM{index:04d}": "제외" for index in range(200)}
    mark(path, wanted)

    parsed = read_bulk_sheet(path)
    assert parsed["counts"]["EXCLUDE"] == 200
    assert parsed["counts"]["blank"] == ASSETS - 200

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        result = apply_bulk(repo, "ITSM", parsed["items"], reason="실물 없음", updated_by="admin")
    assert result["applied"]["EXCLUDE"] == 200

    # 그리고 그 수가 모든 화면에 반영된다.
    with manager.connect() as conn:
        status = ServerStatusService(config, AssetRepository(conn)).status(snapshot_id)
    assert status["all"]["table"]["rows"]["계"]["소계"] == ASSETS - 200


def test_a_mixed_sheet_does_each_row_as_written(portal, tmp_path):
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    mark(path, {
        "CM0000": "제외", "CM0001": "제외",
        "CM0002": "포함",
        "CM0003": "자동",
        "CM0004": "Y",            # 빨리 쓰는 표기도 받는다
    })
    parsed = read_bulk_sheet(path)
    assert parsed["counts"]["EXCLUDE"] == 3
    assert parsed["counts"]["INCLUDE"] == 1
    assert parsed["counts"]["AUTO"] == 1


def test_a_per_row_reason_is_kept(portal, tmp_path):
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    mark(path, {"CM0000": "제외", "CM0001": "제외"},
         reasons={"CM0000": "테스트 장비", "CM0001": "반납 예정"})
    parsed = read_bulk_sheet(path)
    with manager.connect() as conn:
        repo = AssetRepository(conn)
        apply_bulk(repo, "ITSM", parsed["items"], reason="기본사유", updated_by="admin")
        rules = {row["asset_key"]: row["reason"] for row in list_rules(repo, "ITSM")}
    assert rules["CM0000"] == "테스트 장비"
    assert rules["CM0001"] == "반납 예정"


def test_an_empty_action_column_takes_the_mode_from_the_screen(portal, tmp_path):
    """엑셀에서 걸러 남긴 목록을 그대로 올리는 쪽이 빠른 경우가 많다."""
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    parsed = read_bulk_sheet(path, default_mode="EXCLUDE")
    assert parsed["counts"]["EXCLUDE"] == ASSETS
    assert parsed["counts"]["blank"] == 0


# ── 조용히 넘기지 않는가 ─────────────────────────────────────────────────
def test_an_unknown_word_is_reported_not_ignored(portal, tmp_path):
    """300 줄 중 몇 줄이 오타라서 빠졌다는 것을 모르면 숫자가 틀린 채로 보고가 나간다."""
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    mark(path, {"CM0000": "제외", "CM0001": "빼기", "CM0002": "ㅈㅇ"})
    parsed = read_bulk_sheet(path)
    assert parsed["counts"]["EXCLUDE"] == 1
    assert parsed["counts"]["unknown"] == 2
    values = {item["value"] for item in parsed["unknown"]}
    assert values == {"빼기", "ㅈㅇ"}
    # 몇 행인지 알려 줘야 고칠 수 있다.
    assert all(item["row"] > 5 for item in parsed["unknown"])


def test_the_same_asset_twice_uses_the_first_row(portal, tmp_path):
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    workbook = load_workbook(path)
    sheet = workbook.active
    key_at = key_column(sheet)
    sheet.cell(row=6, column=1, value="제외")
    # 같은 자산번호를 아래에 한 번 더 적는다. 사람이 복사하다 흔히 생긴다.
    extra = sheet.max_row + 1
    sheet.cell(row=extra, column=key_at, value=sheet.cell(row=6, column=key_at).value)
    sheet.cell(row=extra, column=1, value="포함")
    workbook.save(path)
    workbook.close()

    parsed = read_bulk_sheet(path)
    assert parsed["counts"]["EXCLUDE"] == 1
    assert parsed["counts"]["INCLUDE"] == 0
    assert parsed["counts"]["duplicated"] == 1


def test_a_sheet_without_an_asset_number_column_is_refused(tmp_path):
    path = tmp_path / "wrong.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        csv.writer(stream).writerows([["이름", "메모"], ["서버1", "제외"]])
    with pytest.raises(AssetScopeError) as excinfo:
        read_bulk_sheet(path)
    assert "자산번호" in str(excinfo.value)


def test_a_plain_csv_made_by_hand_also_works(tmp_path):
    """엑셀을 거치지 않고 메모장으로 만든 목록도 받는다."""
    path = tmp_path / "list.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["CM_ID", "처리", "사유"])
        writer.writerow(["CM0001", "제외", "실물 없음"])
        writer.writerow(["CM0002", "제외", ""])
    parsed = read_bulk_sheet(path)
    assert parsed["counts"]["EXCLUDE"] == 2
    assert parsed["items"][0]["reason"] == "실물 없음"


def test_an_asset_number_that_is_only_digits_is_not_turned_into_a_float(tmp_path):
    """엑셀이 숫자로 바꿔 둔 자산번호가 '1234.0' 이 되면 안 된다."""
    from openpyxl import Workbook

    path = tmp_path / "numeric.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CM_ID", "처리"])
    sheet.append([1234, "제외"])
    workbook.save(path)
    workbook.close()

    parsed = read_bulk_sheet(path)
    assert parsed["items"][0]["asset_key"] == "1234"


def test_the_header_can_sit_on_any_row(tmp_path):
    """우리가 내보낸 파일은 머리글이 4·5행이고, 직접 만든 파일은 1행이다."""
    from openpyxl import Workbook

    path = tmp_path / "shifted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["자산 목록"])
    sheet.append([])
    sheet.append(["설명 줄"])
    sheet.append(["처리", "자산번호"])
    sheet.append(["제외", "CM0009"])
    workbook.save(path)
    workbook.close()

    parsed = read_bulk_sheet(path)
    assert parsed["header_row"] == 4
    assert parsed["items"] == [{"asset_key": "CM0009", "mode": "EXCLUDE", "reason": ""}]


def test_the_fill_in_column_and_the_current_state_have_different_names(portal, tmp_path):
    """사람이 채우는 칸과 지금 상태가 같은 이름이면 어느 칸인지 헷갈린다.

    읽는 쪽도 머리글 이름으로 열을 찾으므로, 같은 이름이 둘 있으면 엉뚱한 칸을
    읽는다.
    """
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    labels = [cell.value for cell in load_workbook(path).active[4]]
    assert labels.count("제외 사유") == 1, labels
    assert "제외 사유(현재)" in labels


def test_what_is_already_excluded_is_marked_in_the_file(portal, tmp_path):
    """이미 수동으로 제외해 둔 것은 파일에 표시돼야 두 번 적지 않는다."""
    from asset_sync.services.asset_scope import save_rules

    config, manager, snapshot_id = portal
    with manager.connect() as conn:
        save_rules(AssetRepository(conn), "ITSM", [{"asset_key": "CM0005"}],
                   mode="EXCLUDE", reason="먼저 뺀 것")
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    sheet = load_workbook(path).active
    key_at = key_column(sheet)
    for line in range(6, sheet.max_row + 1):
        if sheet.cell(row=line, column=key_at).value == "CM0005":
            assert sheet.cell(row=line, column=1).value == "제외"
            assert sheet.cell(row=line, column=2).value == "먼저 뺀 것"
            break
    else:
        pytest.fail("CM0005 를 찾지 못했습니다")


# ── 엑셀에서 값 자체를 고쳐 올리기 ───────────────────────────────────────
def _raw_column(sheet, name: str) -> int:
    return [cell.value for cell in sheet[5]].index(name) + 1


def _label_column(sheet, name: str) -> int:
    return [cell.value for cell in sheet[4]].index(name) + 1


def test_the_computed_columns_are_locked_in_excel(portal, tmp_path):
    """계산값은 고쳐도 반영되지 않는다. 그러니 애초에 못 고치게 잠근다.

    고친 뒤 "왜 반영이 안 되지" 를 묻는 것보다 엑셀이 그 자리에서 막는 쪽이 낫다.
    """
    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    sheet = load_workbook(path).active

    assert sheet.protection.sheet is True, "시트가 잠겨 있지 않습니다"
    # 고쳐야 하는 칸은 열려 있어야 한다.
    for name in ("처리", "제외 사유"):
        assert sheet.cell(row=6, column=_label_column(sheet, name)).protection.locked is False, name
    assert sheet.cell(row=6, column=_raw_column(sheet, "CM_EOL_DT")).protection.locked is False
    # 계산값 칸은 잠겨 있어야 한다.
    assert sheet.cell(row=6, column=_label_column(sheet, "EOSL 연도")).protection.locked is True
    # 잠갔더라도 정렬·필터는 되어야 한다. 막으면 쓸 수가 없다.
    assert sheet.protection.autoFilter is False
    assert sheet.protection.sort is False


def test_filling_in_a_raw_column_becomes_a_correction(portal, tmp_path):
    """ITSM 에 EOSL 이 없어 '미사용' 으로 잡히는 자산을 엑셀에서 채워 넣는다."""
    from asset_sync.services.asset_scope import apply_corrections, current_raw_values

    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")

    with manager.connect() as conn:
        current = current_raw_values(config, AssetRepository(conn), snapshot_id)

    workbook = load_workbook(path)
    sheet = workbook.active
    column = _raw_column(sheet, "CM_EOL_DT")
    keys = []
    for line in range(6, 9):
        sheet.cell(row=line, column=column, value="2032-06-30")
        keys.append(sheet.cell(row=line, column=key_column(sheet)).value)
    workbook.save(path)
    workbook.close()

    parsed = read_bulk_sheet(path, current=current)
    assert parsed["counts"]["corrections"] == 3
    assert {item["field"] for item in parsed["corrections"]} == {"CM_EOL_DT"}
    assert {item["new"] for item in parsed["corrections"]} == {"2032-06-30"}
    # 제외는 하지 않았다. 값만 고친 것이다.
    assert parsed["counts"]["total"] == 0

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        result = apply_corrections(repo, parsed["corrections"], reason="담당자 확인", updated_by="admin")
        assert result["applied"] == 3
        # 보정은 수집 결과를 건드리지 않는다. 다시 수집해도 남아야 한다.
        after = current_raw_values(config, repo, snapshot_id)
    for key in keys:
        assert after[key]["CM_EOL_DT"] == "2032-06-30"


def test_the_same_file_twice_does_not_pile_up_corrections(portal, tmp_path):
    """두 번째 업로드에서 같은 보정이 또 쌓이면 어느 것이 적용되는지 알 수 없다."""
    from asset_sync.services.asset_scope import apply_corrections, current_raw_values

    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=6, column=_raw_column(sheet, "CM_EOL_DT"), value="2032-06-30")
    workbook.save(path)
    workbook.close()

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        first = read_bulk_sheet(path, current=current_raw_values(config, repo, snapshot_id))
        apply_corrections(repo, first["corrections"], updated_by="admin")

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        second = read_bulk_sheet(path, current=current_raw_values(config, repo, snapshot_id))
        assert second["counts"]["corrections"] == 0
        rows = conn.execute("SELECT COUNT(*) AS n FROM manual_asset_override").fetchone()
        assert dict(rows)["n"] == 1


def test_clearing_a_corrected_value_takes_the_correction_back(portal, tmp_path):
    from asset_sync.services.asset_scope import apply_corrections, current_raw_values

    config, manager, snapshot_id = portal
    path = export_list(config, manager, snapshot_id, tmp_path / "out")
    workbook = load_workbook(path)
    sheet = workbook.active
    column = _raw_column(sheet, "CM_EOL_DT")
    sheet.cell(row=6, column=column, value="2032-06-30")
    workbook.save(path)
    workbook.close()

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        apply_corrections(
            repo, read_bulk_sheet(path, current=current_raw_values(config, repo, snapshot_id))["corrections"],
            updated_by="admin")

    # 이제 그 칸을 비워 다시 올리면 원래 값으로 돌아간다.
    workbook = load_workbook(path)
    sheet = workbook.active
    # openpyxl 의 cell(value=None) 은 아무 일도 하지 않는다. 값을 직접 지운다.
    sheet.cell(row=6, column=column).value = None
    workbook.save(path)
    workbook.close()

    with manager.connect() as conn:
        repo = AssetRepository(conn)
        parsed = read_bulk_sheet(path, current=current_raw_values(config, repo, snapshot_id))
        assert parsed["counts"]["corrections"] == 1
        result = apply_corrections(repo, parsed["corrections"], updated_by="admin")
        assert result["cleared"] == 1
        left = conn.execute("SELECT COUNT(*) AS n FROM manual_asset_override").fetchone()
        assert dict(left)["n"] == 0


def test_a_hand_made_sheet_does_not_trigger_corrections(tmp_path):
    """머리글이 한 줄이면 어느 열이 원본 컬럼인지 알 수 없다. 건드리지 않는다."""
    path = tmp_path / "hand.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["CM_ID", "처리"])
        writer.writerow(["CM0001", "제외"])
    parsed = read_bulk_sheet(path, current={"CM0001": {"CM_EOL_DT": "2030-12-31"}})
    assert parsed["counts"]["corrections"] == 0
    assert parsed["counts"]["EXCLUDE"] == 1
