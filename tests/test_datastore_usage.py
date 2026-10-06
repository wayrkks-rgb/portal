"""데이터스토어 디스크는 '쓴 양' 과 '나눠준 양' 이 다르다.

사용률  = (용량 − 여유) ÷ 용량            → 지금 얼마나 차 있나
할당률  = (사용 + 미사용 씬 몫) ÷ 용량     → VM 에게 얼마를 약속했나

씬 프로비저닝이면 할당률이 100% 를 넘을 수 있다(과할당). 사용률만 보고 "아직
반이나 남았다" 고 읽으면, VM 이 약속받은 만큼 채우는 순간 데이터스토어가 꽉 찬다.
그래서 두 값을 따로 내고, 넘긴 것을 따로 표시한다.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from asset_sync.config import AppConfig
from asset_sync.db.sqlite_manager import SQLiteManager
from asset_sync.repositories import AssetRepository
from asset_sync.services.resource_usage_service import VMResourceUsageExportService

STAT_DATE = "2026-09-30"
GB = 1024
TB = 1024 * GB
ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "collect_vcenter_resource_usage.ps1"


def _service(tmp_path: Path):
    cfg = AppConfig(
        root_dir=tmp_path,
        sqlite_path=Path("data/disk.db"),
        rvtools={"resource_usage": {"enabled": False}, "vcenters": []},
    )
    manager = SQLiteManager(cfg.database_path)
    manager.initialize()
    return cfg, manager


def _run(conn, stat_date: str = STAT_DATE) -> int:
    cur = conn.execute(
        "INSERT INTO resource_usage_run(period_start, period_end, started_at, status)"
        " VALUES(?, ?, ?, 'SUCCESS')",
        (stat_date, stat_date, stat_date),
    )
    return int(cur.lastrowid)


def _datastore(conn, run_id: int, name: str, *, capacity_mb: int, free_mb: int,
               provisioned_mb: int, accessible: int = 1, cluster: str | None = None,
               stat_date: str = STAT_DATE, vcenter: str = "VC1") -> None:
    conn.execute(
        "INSERT INTO datastore_usage_daily(run_id, stat_date, vcenter_id, service_name,"
        " cluster_name, datastore_name, datastore_type, accessible, capacity_mb, free_mb,"
        " used_mb, provisioned_mb, host_count, vm_count, collection_status, raw_json, created_at)"
        " VALUES(?,?,?,'업무A',?,?, 'VMFS', ?,?,?,?,?, 4, 0, 'SUCCESS', '{}', ?)",
        (run_id, stat_date, vcenter, cluster, name, accessible, capacity_mb, free_mb,
         capacity_mb - free_mb, provisioned_mb, stat_date),
    )


def _host(conn, run_id: int, esxi: str, *, cluster: str = "CL1", stat_date: str = STAT_DATE) -> None:
    conn.execute(
        "INSERT INTO host_resource_usage_daily(run_id, stat_date, vcenter_id, service_name,"
        " cluster_name, esxi_host, vm_count, allocated_cpu_cores, allocated_memory_mb,"
        " cpu_max_pct, cpu_avg_pct, mem_max_pct, mem_avg_pct, sample_count, collection_status,"
        " raw_json, created_at) VALUES(?,?,'VC1','업무A',?,?, 0, 64, ?, 80,40,70,35, 12,"
        " 'SUCCESS', '{}', ?)",
        (run_id, stat_date, cluster, esxi, 512 * GB, stat_date),
    )


def _vm(conn, run_id: int, name: str, *, esxi: str, provisioned_mb: int, used_mb: int,
        cluster: str = "CL1", stat_date: str = STAT_DATE) -> None:
    conn.execute(
        "INSERT INTO vm_resource_usage_daily(run_id, stat_date, vcenter_snapshot_id, asset_key,"
        " vcenter_id, service_name, cluster_name, esxi_host, vm_uuid, vm_name, power_state,"
        " allocated_cpu_cores, allocated_memory_mb, provisioned_disk_mb, used_disk_mb,"
        " cpu_max_pct, cpu_avg_pct, mem_max_pct, mem_avg_pct, sample_count, inventory_status,"
        " collection_status, raw_json, created_at)"
        " VALUES(?,?,NULL,?, 'VC1','업무A',?,?,?,?, 'poweredOn', 4, ?, ?, ?,"
        " 30,10,60,30, 12, 'CURRENT', 'SUCCESS', '{}', ?)",
        (run_id, stat_date, name, cluster, esxi, name, name, 16 * GB,
         provisioned_mb, used_mb, stat_date),
    )


def test_used_and_provisioned_are_two_different_numbers(tmp_path: Path) -> None:
    """10TB 중 3TB 가 차 있는데 VM 에게 12TB 를 약속했다. 둘 다 보여야 한다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _datastore(conn, run, "DS_THIN",
                   capacity_mb=10 * TB, free_mb=7 * TB, provisioned_mb=12 * TB)
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    row = result["datastores"][0]
    assert row["used_pct"] == 30.0, "사용률은 지금 차 있는 양이다"
    assert row["provision_pct"] == 120.0, "할당률은 약속한 양이다"
    assert row["over_provisioned"] is True
    assert row["capacity_gb"] == 10 * 1024
    assert row["free_gb"] == 7 * 1024


def test_a_full_datastore_is_not_hidden_by_a_low_allocation(tmp_path: Path) -> None:
    """반대 경우도 있다. 약속은 적은데 실제로 꽉 찬 데이터스토어."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _datastore(conn, run, "DS_FULL",
                   capacity_mb=1 * TB, free_mb=50 * GB, provisioned_mb=600 * GB)
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    row = result["datastores"][0]
    assert row["used_pct"] > 95, "거의 꽉 찼는데 사용률이 낮게 나오면 안 된다"
    assert row["provision_pct"] < 60
    assert row["over_provisioned"] is False


def test_an_unreachable_datastore_is_left_out_of_the_total(tmp_path: Path) -> None:
    """접속할 수 없는 데이터스토어의 용량을 더하면 여유가 있다고 잘못 읽힌다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _datastore(conn, run, "DS_OK", capacity_mb=1 * TB, free_mb=500 * GB, provisioned_mb=600 * GB)
        _datastore(conn, run, "DS_DEAD", capacity_mb=9 * TB, free_mb=9 * TB,
                   provisioned_mb=0, accessible=0)
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    disk = result["disk"]
    assert disk["datastore_count"] == 2
    assert disk["inaccessible_count"] == 1
    assert disk["capacity_gb"] == 1024, "접속불가 9TB 는 용량에서 빠진다"
    assert disk["used_pct"] == 51.17, "1TB 중 500GB 여유 → 524GB 사용"


def test_the_latest_day_wins_and_the_peak_is_kept(tmp_path: Path) -> None:
    """디스크는 평균이 쓸모 없다. 지금 모습은 마지막 날 값이고, 최대는 따로 적는다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        first = _run(conn, "2026-09-01")
        _datastore(conn, first, "DS_A", capacity_mb=1 * TB, free_mb=100 * GB,
                   provisioned_mb=900 * GB, stat_date="2026-09-01")
        last = _run(conn, "2026-09-30")
        _datastore(conn, last, "DS_A", capacity_mb=1 * TB, free_mb=600 * GB,
                   provisioned_mb=500 * GB, stat_date="2026-09-30")
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(
            "2026-09-01", "2026-09-30"
        )

    row = result["datastores"][0]
    assert row["latest_stat_date"] == "2026-09-30"
    assert row["used_pct"] == 41.41, "마지막 날의 모습"
    assert row["used_pct_max"] == 90.23, "기간 중 가장 찼을 때"


def test_vm_disk_rolls_up_to_the_esxi_and_the_cluster(tmp_path: Path) -> None:
    """ESXi 줄에는 '그 위 VM 이 차지한 디스크' 가 보여야 한다.

    디스크 용량은 ESXi 가 아니라 데이터스토어에서 나가므로 ESXi 에는 비율을 쓰지
    않는다. 대신 할당한 것 중 실제로 쓴 몫(실사용률)을 적는다.
    """
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _host(conn, run, "esxi-01")
        _host(conn, run, "esxi-02")
        _vm(conn, run, "vm-a", esxi="esxi-01", provisioned_mb=100 * GB, used_mb=40 * GB)
        _vm(conn, run, "vm-b", esxi="esxi-01", provisioned_mb=100 * GB, used_mb=60 * GB)
        _vm(conn, run, "vm-c", esxi="esxi-02", provisioned_mb=200 * GB, used_mb=50 * GB)
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    hosts = {row["esxi_host"]: row for row in result["hosts"]}
    assert hosts["esxi-01"]["assigned_disk_gb"] == 200
    assert hosts["esxi-01"]["used_disk_gb"] == 100
    assert hosts["esxi-01"]["disk_fill_pct"] == 50.0
    assert hosts["esxi-02"]["assigned_disk_gb"] == 200
    assert hosts["esxi-02"]["disk_fill_pct"] == 25.0

    cluster = result["clusters"][0]
    assert cluster["assigned_disk_gb"] == 400, "클러스터는 ESXi 를 더한 값"
    assert cluster["used_disk_gb"] == 150
    assert cluster["disk_fill_pct"] == 37.5

    vms = {row["vm_name"]: row for row in result["vms"]}
    assert vms["vm-a"]["provisioned_disk_gb"] == 100
    assert vms["vm-a"]["used_disk_gb"] == 40
    assert vms["vm-a"]["disk_fill_pct"] == 40.0


def test_the_excel_has_a_datastore_sheet_with_both_numbers(tmp_path: Path) -> None:
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _datastore(conn, run, "DS_THIN",
                   capacity_mb=10 * TB, free_mb=7 * TB, provisioned_mb=12 * TB)
        conn.commit()
        target = VMResourceUsageExportService(cfg, AssetRepository(conn)).export_xlsx(
            STAT_DATE, STAT_DATE, tmp_path / "out"
        )

    book = load_workbook(target)
    assert "DatastoreUsage" in book.sheetnames
    sheet = book["DatastoreUsage"]
    labels = [cell.value for cell in sheet[1]]
    for label in ("실제 용량 GB", "사용 GB", "사용률 %", "VM 할당(프로비저닝) GB", "할당률 %", "과할당"):
        assert label in labels, f"엑셀에 {label} 칸이 없다"
    row = {label: sheet.cell(2, index).value for index, label in enumerate(labels, start=1)}
    assert row["사용률 %"] == 30.0
    assert row["할당률 %"] == 120.0
    assert row["과할당"] == "과할당"
    # '0.##' 표시형식은 16 을 "16." 으로 그린다. 숫자 칸은 General 이어야 한다.
    assert sheet.cell(2, labels.index("실제 용량 GB") + 1).number_format == "General"
    book.close()


def test_an_old_script_without_datastores_still_collects_hosts_and_vms(tmp_path: Path) -> None:
    """폐쇄망이라 스크립트 반입이 늦을 수 있다. 그 사이에도 나머지는 되어야 한다."""
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        service = VMResourceUsageExportService(cfg, AssetRepository(conn))
        run = _run(conn)
        # 예전 payload: datastores 칸이 아예 없다.
        service._replace_run_rows(run, STAT_DATE, None, [{
            "vcenter_id": "VC1", "service_name": "업무A", "cluster_name": "CL1",
            "esxi_host": "esxi-01", "vm_count": 1,
            "allocated_cpu_cores": 64, "allocated_memory_mb": 512 * GB,
        }], [{
            "vcenter_id": "VC1", "vm_name": "vm-a", "esxi_host": "esxi-01",
            "cluster_name": "CL1", "allocated_cpu_cores": 4, "allocated_memory_mb": 16 * GB,
        }])
        conn.commit()
        result = service.summary(STAT_DATE, STAT_DATE)

    assert len(result["hosts"]) == 1
    assert len(result["vms"]) == 1
    assert result["datastores"] == []
    assert result["disk"]["capacity_gb"] is None or result["disk"]["capacity_gb"] == 0
    assert result["hosts"][0]["assigned_disk_gb"] in (None, 0), "값이 없으면 0 으로 둔다"


# ── 수집 쪽 계약 ────────────────────────────────────────────────────────


def code_only() -> str:
    return "\n".join(
        line for line in SCRIPT.read_text(encoding="utf-8").splitlines()
        if not line.lstrip().startswith("#")
    )


def test_the_script_reads_the_thin_provisioning_number() -> None:
    """Uncommitted 를 안 더하면 씬 디스크의 약속분이 빠져 과할당을 못 본다."""
    text = code_only()
    assert "ViewType Datastore" in text, "데이터스토어를 한 번에 받아야 한다"
    assert "Uncommitted" in text, "씬 프로비저닝 미사용 몫을 더해야 할당량이 된다"
    assert "provisioned_mb" in text and "used_mb" in text and "capacity_mb" in text
    assert "datastores = @($datastoreRows)" in text, "payload 에 데이터스토어가 없다"


def test_the_script_takes_vm_disk_from_the_bulk_view() -> None:
    """$vm.ProvisionedSpaceGB 를 읽으면 VM 하나당 왕복이 또 생긴다."""
    text = code_only()
    assert "Summary.Storage" in text, "디스크는 일괄 조회 속성으로 받아야 한다"
    assert ".ProvisionedSpaceGB" not in text
    assert ".UsedSpaceGB" not in text
    assert "provisioned_disk_mb" in text and "used_disk_mb" in text


def test_the_script_does_not_add_another_round_trip() -> None:
    """Get-View 호출이 늘면 통합기마다 수십 초가 더 걸린다."""
    text = code_only()
    calls = re.findall(r"Get-View -Server \$viServer -ViewType (\w+)", text)
    assert sorted(calls) == ["ClusterComputeResource", "Datastore", "VirtualMachine"], calls


def test_the_screens_show_both_numbers() -> None:
    """화면에 사용률만 있으면 과할당을 못 본다. 양쪽 id 가 맞아야 그려진다."""
    page = (ROOT / "templates" / "pages" / "monthly_check.html").read_text(encoding="utf-8")
    script = (ROOT / "templates" / "partials" / "js" / "monthly_check.html").read_text(encoding="utf-8")
    for element_id in ("monthly-datastore-table", "monthly-datastore-risk",
                       "monthly-disk-capacity", "monthly-disk-used",
                       "monthly-disk-provisioned", "monthly-disk-over"):
        assert f'id="{element_id}"' in page, f"{element_id} 가 화면에 없다"
        assert f"'{element_id}'" in script, f"{element_id} 를 스크립트가 안 찾는다"
    assert "renderMonthlyDatastoresAgain" in script
    assert "과할당" in script, "과할당 표시가 없다"

    report = (ROOT / "templates" / "pages" / "report.html").read_text(encoding="utf-8")
    report_js = (ROOT / "templates" / "partials" / "js" / "dashboard.html").read_text(encoding="utf-8")
    assert 'id="resource-datastore-body"' in report
    assert "resource-datastore-body" in report_js


def test_the_host_vm_count_always_matches_the_vm_list(tmp_path: Path) -> None:
    """통합기 표의 VM 합은 아래 VM 목록의 줄 수와 같아야 한다.

    저장된 수는 수집 당시의 수다. 그 뒤에 VM 을 자산에서 빼면 목록은 줄지만
    저장된 수는 그대로다. 어느 통합기의 VM 이 **전부** 빠지면 예전에는 그
    통합기만 옛 수를 들고 있어서, 표의 합이 목록보다 컸다.
    """
    cfg, manager = _service(tmp_path)
    with manager.connect() as conn:
        run = _run(conn)
        _host(conn, run, "esxi-01")
        _host(conn, run, "esxi-02")
        # 수집 당시 esxi-01 에 3대, esxi-02 에 1대였다고 저장해 둔다.
        conn.execute("UPDATE host_resource_usage_daily SET vm_count=3 WHERE esxi_host='esxi-01'")
        conn.execute("UPDATE host_resource_usage_daily SET vm_count=1 WHERE esxi_host='esxi-02'")
        # 지금 남아 있는 VM 은 esxi-02 의 1대뿐이다(esxi-01 의 3대는 전부 제외됨).
        _vm(conn, run, "vm-live", esxi="esxi-02", provisioned_mb=10 * GB, used_mb=5 * GB)
        conn.commit()
        result = VMResourceUsageExportService(cfg, AssetRepository(conn)).summary(STAT_DATE, STAT_DATE)

    hosts = {row["esxi_host"]: row for row in result["hosts"]}
    assert hosts["esxi-01"]["vm_count"] == 0, "VM 이 전부 빠진 통합기는 0 이어야 한다"
    assert hosts["esxi-01"]["stored_vm_count"] == 3, "저장된 수는 따져볼 수 있게 남긴다"
    assert hosts["esxi-02"]["vm_count"] == 1
    assert sum(h["vm_count"] for h in result["hosts"]) == len(result["vms"])
    # 클러스터 합계도 같은 수를 써야 한다.
    assert sum(c["vm_count"] for c in result["clusters"]) == len(result["vms"])
