"""늘거나 빠진 서버가 **어느 서버인지** 보여야 한다.

"8대 늘었다" 만 알면 그걸 자산에서 빼야 하는지 둬야 하는지 판단할 수 없다. 변동
내역에 호스트명·IP·업무명·OS·EOSL 이 함께 나와야 하고, 표의 증감 숫자를 눌러
그 칸의 서버만 따로 볼 수 있어야 한다. 그 성질을 여기서 지킨다.
"""

from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asset_sync.config import AppConfig
from asset_sync.db.manager import create_manager
from asset_sync.repositories import AssetRepository
from asset_sync.services.server_status_service import ServerStatusService

ROOT = Path(__file__).resolve().parents[1]
THIS_MONTH = date(2026, 9, 30)
LAST_MONTH = date(2026, 8, 31)


def asset(cm_id: str, *, physical: bool = False, place: str = "CMPLACE010",
          os_family: str = "Linux Redhat", eosl: str = "2030-12-31",
          name: str | None = None) -> dict:
    return {
        "cm_id": cm_id,
        "normalized_hostname": f"host-{cm_id}",
        "primary_ip": f"10.0.0.{int(cm_id[-3:]) % 250 + 1}",
        "os_family": os_family,
        "os_version": "8.6",
        "status_code": "CMSTA010",
        "server_category_code": "CMSVRCATCD010" if physical else "CMSVRCATCD020",
        "raw": {
            "CM_ID": cm_id, "CM_NAME": name or f"업무-{cm_id}",
            "CM_HOSTNAME": f"host-{cm_id}", "CM_IP": f"10.0.0.{int(cm_id[-3:]) % 250 + 1}",
            "CM_OS": "CMCIOSCD010", "CM_OS_VERSION": "8.6",
            "CM_EOL_DT": eosl, "CM_PLACE": place,
        },
    }


@pytest.fixture()
def portal(tmp_path: Path):
    config = AppConfig(root_dir=tmp_path, sqlite_path=Path("data/movement.db"))
    manager = create_manager(config)
    manager.initialize()
    return config, manager


def seed(manager, day: date, records: list[dict]) -> int:
    with manager.connect() as conn:
        repo = AssetRepository(conn)
        run = repo.start_collection_run("ITSM", day.isoformat())
        repo.finish_collection_run(run, "SUCCESS", len(records), day.isoformat())
        snapshot_id = repo.create_snapshot(
            "ITSM", day.isoformat(), f"{day.isoformat()}T07:00:00", run, "SUCCESS", len(records), "h"
        )
        conn.executemany(
            "INSERT INTO itsm_asset_snapshot(snapshot_id,cm_id,normalized_hostname,primary_ip,"
            "ip_json,cpu_cores,memory_mb,os_family,os_version,status_code,server_category_code,"
            "environment_code,eos_value,record_hash,raw_json) VALUES(?,?,?,?,'[]',?,?,?,?,?,?,?,?,?,?)",
            [(snapshot_id, r["cm_id"], r["normalized_hostname"], r["primary_ip"], 4, 8192,
              r["os_family"], r["os_version"], r["status_code"], r["server_category_code"],
              "CMOWNCATCD0010", r["raw"].get("CM_EOL_DT"), "h",
              json.dumps(r["raw"], ensure_ascii=False)) for r in records],
        )
        conn.commit()
    return snapshot_id


def test_the_movement_list_says_which_server_with_enough_to_decide(portal) -> None:
    """빼야 할지 둬야 할지 판단할 수 있는 칸이 다 있어야 한다."""
    config, manager = portal
    before = seed(manager, LAST_MONTH, [asset("CM001"), asset("CM002")])
    now = seed(manager, THIS_MONTH, [
        asset("CM001"),
        asset("CM003", physical=True, place="CMPLACE020", name="신규업무"),
    ])

    with manager.connect() as conn:
        service = ServerStatusService(config, AssetRepository(conn))
        movements = service.movements(now, before)

    created = movements["created"]["items"]
    removed = movements["removed"]["items"]
    assert [item["cm_id"] for item in created] == ["CM003"]
    assert [item["cm_id"] for item in removed] == ["CM002"]

    item = created[0]
    # 이 칸이 다 있어야 "이게 뭐지?" 를 사람이 판단할 수 있다.
    assert item["hostname"] == "host-CM003"
    assert item["primary_ip"].startswith("10.0.0.")
    assert item["service_name"] == "신규업무"
    assert item["location"] == "DR"
    assert item["kind"] == "물리서버"
    assert item["os_group"] == "Linux"
    assert item["eosl_year"] == 2030
    assert item["status_label"] == "운영"
    # 대수는 그대로 맞아야 한다. 목록과 숫자가 어긋나면 안 된다.
    assert movements["created"]["total"] == len(created)
    assert movements["removed"]["total"] == len(removed)


def test_the_list_is_sorted_so_it_reads_in_the_same_order_as_the_table(portal) -> None:
    config, manager = portal
    before = seed(manager, LAST_MONTH, [])
    now = seed(manager, THIS_MONTH, [
        asset("CM009", place="CMPLACE020"),
        asset("CM001", physical=True),
        asset("CM005"),
    ])
    with manager.connect() as conn:
        movements = ServerStatusService(config, AssetRepository(conn)).movements(now, before)
    order = [(i["location"], i["kind"], i["hostname"]) for i in movements["created"]["items"]]
    assert order == sorted(order), "위치 → 구분 → 호스트명 순서로 읽혀야 한다"


def test_the_eosl_table_shows_what_changed_per_year(portal) -> None:
    """서버현황만 증감이 보이고 EOSL 은 안 보이면, 늘어난 서버가 어느 연도에
    걸렸는지 알 수 없다."""
    config, manager = portal
    before = seed(manager, LAST_MONTH, [asset("CM001", eosl="2030-12-31")])
    now = seed(manager, THIS_MONTH, [
        asset("CM001", eosl="2030-12-31"),
        asset("CM002", eosl="2030-06-30"),
        asset("CM003", physical=True, eosl="2026-12-31"),
    ])

    with manager.connect() as conn:
        service = ServerStatusService(config, AssetRepository(conn))
        eosl = service.eosl(now, today=THIS_MONTH, previous_snapshot_id=before)

    delta = eosl["delta"]
    assert delta["all"]["전체"] == 2, "전체가 2대 늘었다"
    assert delta["physical"]["전체"] == 1
    # 늘어난 2대가 어느 칸에 들어갔는지 보여야 한다.
    changed = {column: value for column, value in delta["all"].items() if value}
    assert changed, "어느 연도가 늘었는지 비어 있으면 안 된다"
    assert sum(v for k, v in delta["all"].items() if k != "전체") == 2

    # 전월이 없으면 증감은 아예 내지 않는다. 0 으로 적으면 '안 바뀐 것' 으로 읽힌다.
    with manager.connect() as conn:
        plain = ServerStatusService(config, AssetRepository(conn)).eosl(now, today=THIS_MONTH)
    assert plain["delta"] == {}


def test_the_excel_lists_the_changed_servers_not_just_the_counts(portal) -> None:
    config, manager = portal
    before = seed(manager, LAST_MONTH, [asset("CM001"), asset("CM002")])
    now = seed(manager, THIS_MONTH, [asset("CM001"), asset("CM003", name="신규업무")])

    from asset_sync.services.monthly_export_service import MonthlyCheckExportService

    with manager.connect() as conn:
        service = ServerStatusService(config, AssetRepository(conn))
        status = service.status(now, before)
        status["movements"] = service.movements(now, before)
        status["eosl"] = service.eosl(now, today=THIS_MONTH, previous_snapshot_id=before)
        status["as_of"] = THIS_MONTH.isoformat()
        workbook = MonthlyCheckExportService(status, service.records(now)).build("movements")
    target = Path(config.root_dir) / "movements.xlsx"
    workbook.save(target)
    workbook.close()

    sheet = load_workbook(target)["자산 변동 내역"]
    text = "\n".join(
        " ".join(str(cell) for cell in row if cell is not None)
        for row in sheet.iter_rows(values_only=True)
    )
    assert "변동 대상 목록" in text
    assert "host-CM003" in text, "늘어난 서버의 호스트명이 파일에 없다"
    assert "신규업무" in text
    assert "host-CM002" in text, "빠진 서버도 적혀야 한다"


def test_the_screen_can_drill_into_the_delta() -> None:
    """증감 숫자가 눌러지지 않으면 어느 서버인지 끝까지 알 수 없다."""
    script = (ROOT / "templates" / "partials" / "js" / "monthly_check.html").read_text(encoding="utf-8")
    page = (ROOT / "templates" / "pages" / "monthly_check.html").read_text(encoding="utf-8")

    assert "function deltaLink(" in script, "증감을 누를 수 있게 하는 함수가 없다"
    # 서버현황 표와 EOSL 표 둘 다에서 써야 한다.
    status_block = script.split("function renderStatusTable", 1)[1].split("\nfunction ", 1)[0]
    eosl_block = script.split("function renderEosl(", 1)[1].split("\nfunction ", 1)[0]
    assert "deltaLink(" in status_block, "서버현황 표의 증감이 안 눌러진다"
    assert "deltaLink(" in eosl_block, "EOSL 표의 증감이 안 눌러진다"
    assert "openMovementBrowser" in script

    # 변동 내역 목록과 검색, 그 자리에서의 제외까지.
    for element_id in ("monthly-movement-items", "monthly-movement-search", "monthly-movement-shown"):
        assert f'id="{element_id}"' in page, f"{element_id} 가 화면에 없다"
        assert f"'{element_id}'" in script, f"{element_id} 를 스크립트가 안 찾는다"
    assert "excludeOneAsset" in script and "excludeOneAsset" in script
    assert "renderMovementItems" in script


def test_the_audit_panel_is_wired() -> None:
    script = (ROOT / "templates" / "partials" / "js" / "monthly_check.html").read_text(encoding="utf-8")
    page = (ROOT / "templates" / "pages" / "monthly_check.html").read_text(encoding="utf-8")
    for element_id in ("monthly-audit-verdict", "monthly-audit-table",
                       "monthly-audit-sources", "monthly-audit-detail"):
        assert f'id="{element_id}"' in page, f"{element_id} 가 화면에 없다"
        assert f"'{element_id}'" in script, f"{element_id} 를 스크립트가 안 찾는다"
    assert "/api/asset-sync/count-audit" in script


# ── 증감 숫자와 목록의 줄 수가 맞아야 한다 ──────────────────────────────


@pytest.fixture()
def web(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    (tmp_path / "config").mkdir()
    (tmp_path / "config" / "app_config.yaml").write_text("itsm:\n  collection_mode: DEMO\n", encoding="utf-8")
    monkeypatch.setenv("ASSET_APP_ROOT", str(tmp_path))
    monkeypatch.setenv("FLASK_SECRET_KEY", "test-secret")
    monkeypatch.chdir(tmp_path)
    from application import create_app
    from application.db import reset_database_manager

    reset_database_manager()
    yield create_app()
    reset_database_manager()


def admin_client(app):
    client = app.test_client()
    with client.session_transaction() as session:
        session["user"] = {"id": 1, "username": "admin", "role": "admin", "name": "admin"}
    return client


def test_clicking_the_delta_returns_exactly_those_servers(web) -> None:
    """표의 (+2) 를 누르면 2줄이 나와야 한다. 숫자와 목록이 어긋나면 안 된다."""
    from asset_sync.config import load_config
    from asset_sync.db.manager import create_manager

    config = load_config()
    manager = create_manager(config)
    seed(manager, LAST_MONTH, [asset("CM001"), asset("CM002")])
    seed(manager, THIS_MONTH, [
        asset("CM001"),
        asset("CM003", place="CMPLACE020"),
        asset("CM004", physical=True),
    ])
    client = admin_client(web)
    month = THIS_MONTH.strftime("%Y-%m")

    status = client.get(f"/api/asset-sync/server-status?month={month}").get_json()
    total_delta = status["all"]["delta"]["계"]["소계"]
    assert total_delta == 1, "2대 늘고 1대 빠졌으니 순증감은 +1"

    created = client.get(f"/api/asset-sync/assets?month={month}&change=created").get_json()
    assert sorted(item["cm_id"] for item in created["items"]) == ["CM003", "CM004"]
    assert created["change"] == "created"
    assert "신규" in created["change_note"]

    removed = client.get(f"/api/asset-sync/assets?month={month}&change=removed").get_json()
    assert [item["cm_id"] for item in removed["items"]] == ["CM002"]
    assert "삭제" in removed["change_note"]

    # 칸을 좁히면 그 칸의 증감만 나온다. DR 에 1대 늘었다.
    dr = client.get(f"/api/asset-sync/assets?month={month}&change=created&location=DR").get_json()
    assert [item["cm_id"] for item in dr["items"]] == ["CM003"]
    assert status["all"]["delta"]["DR"]["소계"] == len(dr["items"])

    # 물리 표의 증감도 같아야 한다.
    physical = client.get(
        f"/api/asset-sync/assets?month={month}&change=created&kind=physical"
    ).get_json()
    assert [item["cm_id"] for item in physical["items"]] == ["CM004"]
    assert status["physical"]["delta"]["계"]["소계"] == len(physical["items"])


def test_without_a_previous_month_the_delta_list_says_so(web) -> None:
    from asset_sync.config import load_config
    from asset_sync.db.manager import create_manager

    config = load_config()
    manager = create_manager(config)
    seed(manager, THIS_MONTH, [asset("CM001")])
    client = admin_client(web)
    month = THIS_MONTH.strftime("%Y-%m")

    created = client.get(f"/api/asset-sync/assets?month={month}&change=created").get_json()
    assert created["items"] == []
    assert "전월" in created["change_note"]
